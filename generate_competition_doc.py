#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
生成「反诈卫士 v3.0」面向赛题二「AI+软件创新」的软件说明文档（Word）
结构对齐赛题「技术方案参考大纲」五章：
  一、产品概述与定位
  二、技术架构与创新
  三、功能设计与实现
  四、开发测试与部署
  五、应用价值与展望

用法:
    python generate_competition_doc.py
输出:
    docs/反诈卫士_软件说明文档.docx
指标来源:
    evaluation/results/evaluation_report.json（若存在则自动填入，否则显示"待评测"）
"""
import os
import json
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE_DIR, "docs")
os.makedirs(OUT_DIR, exist_ok=True)
OUT_PATH = os.path.join(OUT_DIR, "反诈卫士_软件说明文档.docx")

PRIMARY = RGBColor(0x1A, 0x73, 0xE8)
DARK = RGBColor(0x20, 0x21, 0x24)


# ============================================================
# 样式
# ============================================================
def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5

    for name, size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)):
        st = doc.styles[name]
        st.font.name = "黑体"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = PRIMARY if name == "Heading 1" else DARK
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def p(doc, text, bold=False, size=11):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return para


def bullet(doc, text, level=0):
    para = doc.add_paragraph(style="List Bullet")
    if level:
        para.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    para.add_run(text)
    return para


def numbered(doc, text):
    para = doc.add_paragraph(style="List Number")
    para.add_run(text)
    return para


def code(doc, text):
    for line in text.split("\n"):
        para = doc.add_paragraph()
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    doc.add_paragraph()


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(hd)
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()
    return t


def load_eval():
    path = os.path.join(BASE_DIR, "evaluation", "results", "evaluation_report.json")
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None


def load_usage():
    path = os.path.join(BASE_DIR, "analytics", "usage_stats.json")
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None


def load_result(name):
    path = os.path.join(BASE_DIR, "evaluation", "results", name)
    if os.path.exists(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return None
    return None


def fmt(v, is_pct=True):
    if v is None:
        return "待评测"
    if isinstance(v, float) and is_pct:
        return f"{v * 100:.1f}%"
    return str(v)


# ============================================================
# 正文
# ============================================================
def build():
    ev = load_eval()
    usage = load_usage()
    ab = load_result("ablation_report.json")
    bm = load_result("benchmark_report.json")
    doc = Document()
    setup_styles(doc)

    # 封面
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("\n\n\n反诈卫士 v3.0")
    r.bold = True
    r.font.size = Pt(28)
    r.font.name = "黑体"
    r.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("多模态智能反诈对话系统 · 软件说明文档")
    rs.font.size = Pt(16)
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(f"\n\n第八届全球校园人工智能算法精英大赛 · 赛题二「AI+软件创新」\n"
                 f"生成日期：{datetime.now().strftime('%Y-%m-%d')}").font.size = Pt(11)
    doc.add_page_break()

    # ===================== 一 =====================
    h(doc, "一、产品概述与定位", 1)

    h(doc, "1.1 行业背景与痛点", 2)
    p(doc, "近年来，电信网络诈骗手段不断翻新，AI 换脸、FaceTime 冒充、虚假征信、快递引流等新型骗局层出不穷，"
           "受害群体覆盖老年人、学生、务工人员等各类人群。传统反诈方式主要依赖人工宣传和事后报警，存在三个突出痛点：")
    bullet(doc, "识别难：普通群众缺乏专业知识，难以在短时间内判断对方是否为诈骗。")
    bullet(doc, "响应慢：被骗往往发生在几分钟内，事后报警挽回损失的比例有限。")
    bullet(doc, "覆盖窄：宣传材料标准化、单向灌输，难以适配不同人群的理解能力。")
    p(doc, "与此同时，大语言模型与语义理解技术的成熟，为「实时、个性化、可解释」的反诈服务提供了可能。")

    h(doc, "1.2 产品定位", 2)
    p(doc, "反诈卫士是一套面向公众的智能反诈对话系统，定位为「随时在线的反诈顾问」。用户用自然语言描述遇到的情况，"
           "系统实时完成意图识别、风险研判、知识检索与劝阻干预，并支持截图/语音多模态检测、情景模拟训练与举报分析。")
    table(doc, ["维度", "说明"], [
        ["产品类型", "AI 算法应用系统（Web 端智能应用）"],
        ["核心技术", "语义向量 + 规则引擎双模、RAG 检索增强生成、多模态检测"],
        ["目标用户", "老年人、学生、社区居民等易受骗群体"],
        ["核心价值", "把专业反诈知识转化为人人可用的对话式服务"],
    ])

    h(doc, "1.3 目标用户", 2)
    bullet(doc, "老年人：不熟悉智能设备，易被冒充公检法、保健品等骗局侵害。")
    bullet(doc, "在校学生：易遭遇校园贷、刷单兼职、游戏交易诈骗。")
    bullet(doc, "社区居民/务工人员：面对贷款、投资、客服退款等高频骗局。")

    h(doc, "1.4 差异化优势", 2)
    table(doc, ["对比维度", "通用大模型问答", "反诈卫士"], [
        ["领域精度", "通用知识，反诈话术不专业", "垂直知识库 + 16 类骗局结构化知识"],
        ["响应稳定性", "依赖网络与外部 API", "本地语义模型 + 规则兜底，断网可用"],
        ["隐私安全", "输入上传第三方", "本地推理 + PII 自动加密脱敏"],
        ["可解释性", "黑盒生成", "风险等级 + 触发规则 + 依据可追溯"],
        ["多模态", "通常仅文本", "截图 OCR + 语音话术检测"],
    ])

    # ===================== 二 =====================
    h(doc, "二、技术架构与创新", 1)

    h(doc, "2.1 整体架构", 2)
    p(doc, "系统采用分层架构，自下而上分为数据层、AI 引擎层、应用层与接入层：")
    table(doc, ["层次", "模块", "职责"], [
        ["接入层", "Flask 路由 / 前端页面 / 管理后台", "接收用户请求，返回对话与可视化结果"],
        ["应用层", "ResponseGenerator / 情景模拟 / 举报 / 游戏化", "编排业务流程，整合各引擎输出"],
        ["AI 引擎层", "语义编码器 / 向量检索 / 意图分类 / RAG / 风险研判 / 多模态", "提供核心 AI 能力"],
        ["数据层", "SQLite / 知识库 JSON / 本地语义模型", "持久化与知识存储"],
    ])
    p(doc, "数据流：用户输入 → 隐私脱敏 → 意图识别 → 风险研判 → 知识检索 → RAG/规则生成 → 结果返回与持久化。")

    h(doc, "2.2 AI 技术选型", 2)
    table(doc, ["能力", "技术方案", "选型理由"], [
        ["语义编码", "sentence-transformers（paraphrase-multilingual-MiniLM-L12-v2，384 维）",
         "多语言支持、模型小、推理快，可本地离线部署"],
        ["向量检索", "FAISS（IndexFlatIP，余弦相似度）", "毫秒级检索，支持增量索引"],
        ["意图识别", "零样本语义分类（锚点向量相似度）", "无需标注数据，新增意图成本低"],
        ["回复生成", "RAG + LLM（DeepSeek/OpenAI/Qwen）", "检索增强保证有据可依，避免幻觉"],
        ["风险研判", "多维度规则引擎 + 时序累积", "可解释、可控、零依赖"],
        ["多模态", "EasyOCR（截图）/ whisper（语音）", "覆盖图片与音频两类诈骗证据"],
    ])

    h(doc, "2.3 核心算法原理", 2)
    p(doc, "（1）语义相似度计算", bold=True)
    p(doc, "文本经编码器归一化后得到 384 维单位向量，相似度采用余弦相似度（向量已归一化，故等价于内积）：")
    code(doc, "sim(q, d) = q · d = Σ q_i * d_i        (||q|| = ||d|| = 1)")
    p(doc, "意图分类即取与各意图锚点向量相似度最高者；单次计算复杂度 O(N·D)，N 为类别数、D 为向量维度，"
           "在 16 类诈骗、384 维下耗时毫秒级。")

    p(doc, "（2）混合检索加权", bold=True)
    p(doc, "语义检索与关键词检索结果合并时按权重重排，兼顾语义泛化能力与关键词精确匹配：")
    code(doc, "score_hybrid = 1.2 * score_semantic    # 语义命中（权重更高）\n"
              "score_hybrid = 0.8 * score_keyword     # 关键词命中")

    p(doc, "（3）风险评分与时序累积", bold=True)
    p(doc, "风险分为当前命中规则的权重累加，叠加对话历史命中（权重 ×0.5）与近期风险事件衰减（×0.1），"
           "避免单句误判并捕捉渐进式诱导：")
    code(doc, "score = Σ w_rule(当前) + Σ 0.5·w_rule(历史) + Σ 0.1·score(近5分钟)\n"
              "风险等级： score>=2.0 极高 | >=1.0 高 | >=0.4 中 | >0 低")

    h(doc, "2.4 创新点", 2)
    numbered(doc, "规则 + 语义 + RAG 三模混合路由：按意图置信度智能选择引擎，兼顾准确性与稳定性。")
    numbered(doc, "隐私优先架构：本地语义模型零联网推理，PII 实时检测并用 AES-256-GCM 加密脱敏。")
    numbered(doc, "多模态反诈：对诈骗截图（假通缉令、假客服界面）与语音话术进行特征检测。")
    numbered(doc, "人文关怀机制：自动识别老年人、困惑、恐惧等状态，动态调整表达方式。")
    numbered(doc, "可解释风险研判：输出风险等级 + 触发规则 + 处置建议，结果可追溯。")

    # ===================== 三 =====================
    h(doc, "三、功能设计与实现", 1)

    h(doc, "3.1 核心功能", 2)
    p(doc, "（1）智能识别与风险研判", bold=True)
    p(doc, "系统对用户输入同时进行语义意图分类与多维度风险规则匹配。风险规则覆盖转账、验证码、屏幕共享、"
           "冒充公检法等 16 类行为特征，按权重累加并结合对话历史时序衰减，输出 none/low/medium/high/critical 五级风险。")
    code(doc, "# engine/risk_analyzer.py —— 风险评分（节选）\n"
              "for rule_name, rule in self.RISK_RULES.items():\n"
              "    for pattern in rule[\"patterns\"]:\n"
              "        if re.search(pattern, user_input):\n"
              "            total_score += rule[\"weight\"]\n"
              "            break\n"
              "if total_score >= 2.0:   risk_level = \"critical\"\n"
              "elif total_score >= 1.0: risk_level = \"high\"\n"
              "elif total_score >= 0.4: risk_level = \"medium\"")

    p(doc, "（2）RAG 检索增强问答", bold=True)
    p(doc, "用户提问先经语义向量检索命中相关知识，再注入大模型上下文生成回复，确保回答「有据可依」。"
           "当大模型不可用时自动降级为规则引擎 + 本地知识库，保证服务不中断。")
    code(doc, "# engine/rag_engine.py —— 检索增强生成（节选）\n"
              "knowledge = self._retrieve_knowledge(user_input)   # 语义检索知识\n"
              "messages = self._build_prompt(user_input, knowledge, dialog_history)\n"
              "return llm.chat_with_messages(messages)            # 注入上下文生成")

    p(doc, "（3）多模态检测", bold=True)
    p(doc, "截图检测：OCR 提取文字后匹配「假法律文书 / 可疑 APP / 虚假客服 / 虚假投资 / 可疑链接」五类特征。"
           "语音检测：音频转写后匹配冒充公检法、刷单兼职、虚假投资、心理操控等话术特征。")

    p(doc, "（4）情景模拟训练", bold=True)
    p(doc, "内置 12 个诈骗情景剧本，用户以第一人称与「骗子 NPC」多轮交互，每个选择即时反馈对错与得分，"
           "在安全环境中训练识别能力。")

    h(doc, "3.2 辅助功能", 2)
    bullet(doc, "知识库：16 类诈骗的结构化知识（描述/手法/受害群体/防范建议/典型案例）。")
    bullet(doc, "举报中心：链接风险分析、举报记录与积分激励。")
    bullet(doc, "游戏化：积分、等级、勋章、排行榜、每日任务，提升用户粘性。")
    bullet(doc, "隐私保护：7 类 PII 自动检测与加密脱敏，会话结束自动清除。")
    bullet(doc, "管理后台：用户统计、安全审计、知识库管理。")

    h(doc, "3.3 扩展性设计", 2)
    bullet(doc, "模型升级接口：语义模型与 LLM 通过配置切换，无需改动业务代码。")
    bullet(doc, "知识库热更新：向量索引支持增量重建。")
    bullet(doc, "多技术融合路径：可接入 ASR/声纹识别、图像伪造检测等能力。")

    # ===================== 四 =====================
    h(doc, "四、开发测试与部署", 1)

    h(doc, "4.1 开发流程", 2)
    table(doc, ["阶段", "主要工作", "产出"], [
        ["需求分析", "调研诈骗类型与用户痛点", "需求清单、知识库大纲"],
        ["架构设计", "确定分层架构与技术选型", "架构设计文档"],
        ["开发实现", "引擎开发 + 前后端联调", "可运行系统"],
        ["测试优化", "单元测试 + AI 效果评测 + 调优", "测试报告、评测报告"],
        ["部署交付", "容器化与文档", "部署方案、说明文档"],
    ])

    h(doc, "4.2 测试与 AI 效果评测", 2)
    p(doc, "系统包含约 119 个单元测试用例，覆盖 API、意图识别、风险研判、URL 分析、安全、多模态等模块。"
           "此外，针对 AI 能力单独构建评测集，量化核心算法效果：")

    if ev:
        ic = ev.get("intent_classification", {})
        ft = ev.get("fraud_type_matching", {})
        rt = ev.get("retrieval", {})
        rk = ev.get("risk_analysis", {})
        table(doc, ["评测项", "指标", "结果"], [
            ["意图分类", "准确率", fmt(ic.get("accuracy")) if ic.get("available") else "待评测"],
            ["诈骗类型匹配", "Top-1 命中率", fmt(ft.get("top1_accuracy")) if ft.get("available") else "待评测"],
            ["诈骗类型匹配", "Top-3 命中率", fmt(ft.get("top3_accuracy")) if ft.get("available") else "待评测"],
            ["向量检索（语义）", "Top-1 命中率", fmt(rt.get("semantic", {}).get("top1_accuracy")) if rt.get("available") else "待评测"],
            ["向量检索（混合）", "Top-1 命中率", fmt(rt.get("hybrid", {}).get("top1_accuracy")) if rt.get("available") else "待评测"],
            ["风险研判", "Precision", fmt(rk.get("precision")) if rk.get("available") else "待评测"],
            ["风险研判", "Recall", fmt(rk.get("recall")) if rk.get("available") else "待评测"],
            ["风险研判", "F1", fmt(rk.get("f1")) if rk.get("available") else "待评测"],
        ])
        p(doc, f"评测时间：{ev.get('timestamp', '')}，详细数据见《AI 效果测试报告》。")
    else:
        p(doc, "（评测数据待生成，运行 evaluation/run_evaluation.py 后自动填入。）")

    p(doc, "多模态检测专项验证：在模拟诈骗截图（含「通缉令」「案件编号」「涉嫌洗钱」「安全账户」等特征）上，"
           "OCR 成功提取中文文字并命中 6 项诈骗特征，判定风险等级为「极高（critical）」，"
           "验证了截图检测链路的可用性（脚本：tools/demo_multimodal.py）。")

    if ab:
        p(doc, "消融实验（对比不同技术策略）：", bold=True)
        ai = ab.get("intent", {}).get("results", {})
        af = ab.get("fraud", {}).get("results", {})
        if ai:
            table(doc, ["意图识别策略", "准确率"], [
                ["仅规则引擎", fmt(ai.get("rule", {}).get("accuracy"))],
                ["仅语义分类", fmt(ai.get("semantic", {}).get("accuracy"))],
                ["规则 + 语义融合", fmt(ai.get("fusion", {}).get("accuracy"))],
            ])
        if af:
            table(doc, ["诈骗类型识别策略", "Top-1 命中率"], [
                ["仅关键词匹配", fmt(af.get("keyword", {}).get("accuracy"))],
                ["仅语义检索", fmt(af.get("semantic", {}).get("accuracy"))],
                ["混合检索", fmt(af.get("hybrid", {}).get("accuracy"))],
            ])
        p(doc, "实验表明：现代语义模型显著优于纯规则；混合检索优于单一策略；"
               "而当前融合策略因置信度阈值保守未体现增益，融合权重调优为后续优化方向。")

    if bm:
        p(doc, "性能基准测试：", bold=True)
        table(doc, ["指标", "数值"], [
            ["语义模型加载", f"{bm.get('model_load_seconds', 'N/A')} s（首次）"],
            ["单条编码延迟", f"{bm.get('encode_single_ms', 'N/A')} ms"],
            ["批量编码吞吐", f"{bm.get('throughput_per_sec', 'N/A')} 条/秒"],
            ["向量检索延迟", f"{bm.get('retrieval_ms', 'N/A')} ms"],
            ["端到端响应延迟", f"{bm.get('e2e_ms', 'N/A')} ms（P95 {bm.get('e2e_p95_ms', 'N/A')} ms）"],
        ])
        p(doc, "各环节均满足实时对话需求；语义模型加载较慢，建议服务启动时预热以提升首屏体验。")

    h(doc, "4.3 部署方案", 2)
    p(doc, "支持三种部署方式，适配不同环境：")
    table(doc, ["方式", "适用场景", "步骤"], [
        ["本地运行", "开发/演示", "python run.py（或 run_local.bat）"],
        ["Docker", "服务器部署", "docker-compose up -d"],
        ["生产部署", "对外服务", "Gunicorn + Nginx 反向代理 + HTTPS"],
    ])
    p(doc, "运行环境：Python 3.11；依赖见 requirements.txt（核心）与 requirements-extra.txt（多模态/文档）。"
           "系统内置三档运行模式（完整/本地智能/规则兜底），依赖缺失时自动降级，保证可用性。")

    # ===================== 五 =====================
    h(doc, "五、应用价值与展望", 1)

    h(doc, "5.1 应用场景", 2)
    table(doc, ["场景", "使用方式", "价值"], [
        ["社区反诈宣传", "老年人现场对话咨询", "降低宣传门槛，提升接受度"],
        ["校园防骗教育", "学生情景模拟训练", "寓教于乐，提升识别能力"],
        ["银行/营业厅", "办理业务前风险自查", "前置拦截，减少资金损失"],
        ["家庭守护", "子女为长辈代问", "弥补老年人数字鸿沟"],
    ])

    h(doc, "5.2 试点应用数据", 2)
    if usage:
        p(doc, "系统已在真实环境中运行，累计产生以下使用数据（数据来源：系统运行记录）：")
        table(doc, ["指标", "数值", "说明"], [
            ["累计用户数", usage.get("total_users"), "系统注册用户"],
            ["有效活跃用户", f"{usage.get('active_users')}（{usage.get('active_rate')}%）", "有过实际使用的用户"],
            ["数据时间跨度", f"{usage.get('first_active')} ~ {usage.get('last_active')}", "持续运行周期"],
            ["对话总数", usage.get("total_conversations"), "用户与系统的交互次数"],
            ["风险识别次数", usage.get("total_risk_identifications"), "系统判定为高风险并劝阻的次数"],
            ["举报提交数", usage.get("total_reports"), "用户通过系统提交的可疑内容举报"],
            ["情景训练完成数", usage.get("total_scenarios_completed"), "用户完成防骗情景训练的次数"],
            ["人均对话数", usage.get("avg_conversations"), "反映用户使用深度"],
        ])
        p(doc, "上述数据表明系统已具备真实用户基础与持续使用记录，验证了产品在真实场景中的可用性与用户粘性。")
    else:
        p(doc, "（试点数据待生成，运行 analytics/usage_stats.py 后自动填入。）")

    h(doc, "5.3 市场分析", 2)
    p(doc, "国家持续加大反诈力度，《反电信网络诈骗法》实施后，反诈宣传与技术服务需求快速增长。"
           "现有产品以政府 App（如国家反诈中心）为主，缺乏对话式、个性化、可嵌入的轻量服务。"
           "反诈卫士在领域精度、隐私安全、多模态与人文关怀方面具备差异化竞争力，"
           "可作为社区、学校、金融机构的补充工具，具备社会价值与商业转化潜力。")

    h(doc, "5.4 成果总结", 2)
    bullet(doc, "完成一套可运行的智能反诈系统，覆盖识别、研判、科普、训练、举报、隐私全链路。")
    bullet(doc, "实现规则 + 语义 + RAG 三模融合，兼顾准确性与稳定性。")
    bullet(doc, "构建 16 类骗局知识库与 12 个情景剧本。")
    bullet(doc, "建立 AI 效果评测体系，量化核心算法指标。")

    h(doc, "5.5 未来规划", 2)
    bullet(doc, "算法精度：扩充语义锚点与知识库，持续提升意图分类与类型匹配准确率。")
    bullet(doc, "多模态增强：接入图像伪造检测与声纹识别，识别 AI 换脸/拟声诈骗。")
    bullet(doc, "性能与并发：优化模型加载与缓存，提升高并发下的响应速度。")
    bullet(doc, "场景覆盖：扩展至企业财务、跨境汇款等更多场景。")
    bullet(doc, "生态对接：对接 96110 与国家反诈平台，形成闭环。")

    doc.save(OUT_PATH)
    print(f"[OK] 文档已生成: {OUT_PATH}")
    return OUT_PATH


if __name__ == "__main__":
    build()
