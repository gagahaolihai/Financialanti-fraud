"""生成项目介绍PPT - 原版(非评分优化版)"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import os

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

PRIMARY = RGBColor(0xC6, 0x28, 0x28)
PRIMARY_DARK = RGBColor(0x8B, 0x00, 0x00)
ACCENT_BLUE = RGBColor(0x15, 0x65, 0xC0)
ACCENT_GOLD = RGBColor(0xF5, 0x7C, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY = RGBColor(0xF0, 0xF0, 0xF0)
LIGHT_RED = RGBColor(0xFF, 0xEB, 0xEE)
DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
PURPLE = RGBColor(0x7B, 0x1F, 0xA2)

FONT_TITLE = '微软雅黑'
FONT_BODY = '微软雅黑'


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=14, bold=False,
                color=BLACK, align=PP_ALIGN.LEFT, font_name=FONT_BODY):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return tf


def add_multiline(slide, left, top, width, height, lines, font_size=13, color=BLACK):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, str):
            text, bold, c = item, False, color
        else:
            text = item[0]
            bold = item[1] if len(item) > 1 else False
            c = item[2] if len(item) > 2 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = c
        p.font.name = FONT_BODY
        p.space_after = Pt(3)
    return tf


def add_rect(slide, left, top, width, height, color, border=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if border:
        shape.line.color.rgb = border
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_circle(slide, left, top, size, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_top_bar(slide, title_text, subtitle_text=None):
    add_rect(slide, 0, 0, 13.333, 1.35, PRIMARY)
    add_rect(slide, 0, 1.35, 13.333, 0.06, ACCENT_GOLD)
    add_textbox(slide, 0.8, 0.15, 12, 0.7, title_text, font_size=30, bold=True, color=WHITE, font_name=FONT_TITLE)
    if subtitle_text:
        add_textbox(slide, 0.8, 0.8, 12, 0.4, subtitle_text, font_size=14, color=RGBColor(0xFF, 0xCD, 0xD2))


def add_footer(slide, page_num):
    add_textbox(slide, 11.5, 7.05, 1.5, 0.4, str(page_num), font_size=10, color=GRAY, align=PP_ALIGN.RIGHT)
    add_textbox(slide, 0.5, 7.05, 4, 0.4, '多模态智能反诈对话系统', font_size=9, color=GRAY)


def add_card(slide, left, top, width, height, title, desc, icon_color=None, tag=None):
    if icon_color is None:
        icon_color = PRIMARY
    add_rect(slide, left, top, width, height, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
    add_rect(slide, left, top, width, 0.06, icon_color)
    add_textbox(slide, left + 0.25, top + 0.15, width - 0.5, 0.35, title, font_size=14, bold=True, color=BLACK)
    add_textbox(slide, left + 0.25, top + 0.5, width - 0.5, height - 0.7, desc, font_size=11, color=GRAY)
    if tag:
        tag_shape = add_rect(slide, left + width - 1.1, top + 0.12, 0.85, 0.28, icon_color)
        tf = tag_shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9)
        p.font.color.rgb = WHITE
        p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.CENTER


# ============================================================
# Slide 1: 封面
# ============================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide1, DARK_BG)
add_rect(slide1, 0, 0, 13.333, 0.15, PRIMARY)
add_rect(slide1, 0, 7.35, 13.333, 0.15, PRIMARY)
add_rect(slide1, 0, 4.2, 13.333, 3.3, RGBColor(0x22, 0x22, 0x3E))
add_rect(slide1, 1.5, 1.5, 0.08, 2.5, ACCENT_GOLD)
add_textbox(slide1, 2.0, 1.6, 10, 1.0, '多模态智能反诈对话系统', font_size=44, bold=True, color=WHITE, font_name=FONT_TITLE)
add_textbox(slide1, 2.0, 2.55, 10, 0.6, 'Multi-modal Intelligent Anti-Fraud Dialogue System', font_size=18, color=RGBColor(0xAA, 0xAA, 0xAA))
add_textbox(slide1, 2.0, 3.4, 10, 0.5, '用AI技术守护每一位用户的财产安全', font_size=20, color=ACCENT_GOLD)
add_textbox(slide1, 2.0, 5.0, 10, 0.4, '睿抗机器人开发者大赛 (RAICOM) 参赛作品', font_size=16, color=RGBColor(0xCC, 0xCC, 0xCC))
add_textbox(slide1, 2.0, 5.5, 10, 0.4, '基于大语言模型 · 多引擎协同 · 游戏化反诈教育', font_size=13, color=GRAY)
add_textbox(slide1, 2.0, 6.3, 10, 0.4, '2026年6月', font_size=13, color=GRAY)

# ============================================================
# Slide 2: 目录
# ============================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide2, WHITE)
add_top_bar(slide2, '项目概览', 'PROJECT OVERVIEW')

items = [
    ('01', '项目背景', '电信诈骗形势严峻，AI反诈势在必行'),
    ('02', '系统架构', '8大引擎协同，多模态智能对话'),
    ('03', '核心功能', '六大功能模块，全方位反诈防护'),
    ('04', '亮点功能', '情景模拟 · 知识库 · 隐私保护'),
    ('05', '技术架构', 'Flask + LLM + 知识图谱 + 游戏化'),
    ('06', '演示视频', '2分36秒完整功能演示'),
]
for i, (num, title, desc) in enumerate(items):
    y = 1.9 + i * 0.9
    add_circle(slide2, 1.2, y, 0.5, PRIMARY if i < 3 else ACCENT_BLUE)
    add_textbox(slide2, 1.2, y + 0.06, 0.5, 0.4, num, font_size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide2, 2.0, y - 0.02, 4, 0.3, title, font_size=18, bold=True, color=BLACK)
    add_textbox(slide2, 2.0, y + 0.32, 8, 0.3, desc, font_size=12, color=GRAY)

add_footer(slide2, 2)

# ============================================================
# Slide 3: 项目背景
# ============================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide3, WHITE)
add_top_bar(slide3, '项目背景与痛点', 'WHY ANTI-FRAUD AI?')

cards = [
    ('📊', '电信诈骗年涉案金额', '超千亿元', '公安部2025年数据'),
    ('👵', '老年群体受骗率', '持续走高', '70岁以上为高危人群'),
    ('⏱️', '诈骗手段迭代周期', '不到30天', '新型骗术层出不穷'),
    ('🛡️', '传统反诈手段', '被动滞后', '缺乏实时交互式防护'),
]
for i, (icon, title, value, sub) in enumerate(cards):
    x = 0.6 + i * 3.15
    add_rect(slide3, x, 2.0, 2.95, 2.8, WHITE, border=RGBColor(0xE8, 0xE8, 0xE8))
    add_rect(slide3, x, 2.0, 2.95, 0.07, PRIMARY)
    add_textbox(slide3, x + 0.3, 2.3, 2.3, 0.5, icon, font_size=28, align=PP_ALIGN.CENTER)
    add_textbox(slide3, x + 0.2, 2.8, 2.5, 0.4, title, font_size=13, color=GRAY, align=PP_ALIGN.CENTER)
    add_textbox(slide3, x + 0.2, 3.15, 2.5, 0.5, value, font_size=24, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_textbox(slide3, x + 0.2, 4.3, 2.5, 0.3, sub, font_size=10, color=GRAY, align=PP_ALIGN.CENTER)

add_multiline(slide3, 0.8, 5.2, 12, 1.5, [
    ('核心痛点：公众反诈知识获取渠道单一，缺乏"即时可触达、场景可模拟、风险可感知"的交互式反诈教育工具', True, PRIMARY),
    ('我们的方案：构建基于多模态大语言模型的智能反诈对话系统，集知识普及、风险识别、情景训练、举报联动于一体', False, ACCENT_BLUE),
], font_size=14)

add_footer(slide3, 3)

# ============================================================
# Slide 4: 系统架构总览
# ============================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide4, WHITE)
add_top_bar(slide4, '系统架构总览', 'SYSTEM ARCHITECTURE')

add_rect(slide4, 0.5, 2.0, 12.3, 0.8, LIGHT_RED)
add_textbox(slide4, 0.5, 2.15, 12.3, 0.5, '用户交互层', font_size=14, bold=True, color=PRIMARY, align=PP_ALIGN.CENTER)
add_textbox(slide4, 0.5, 2.5, 12.3, 0.4, 'Web Chat ｜ 移动端 ｜ API接口', font_size=11, color=GRAY, align=PP_ALIGN.CENTER)

engines = [
    ('意图识别', 'intent'),
    ('风险分析', 'risk'),
    ('知识检索', 'knowledge'),
    ('隐私守护', 'privacy'),
    ('URL分析', 'url'),
    ('LLM客户端', 'llm'),
    ('响应生成', 'response'),
    ('知识图谱', 'graph'),
]
for i, (name, _) in enumerate(engines):
    x = 0.4 + i * 1.58
    add_rect(slide4, x, 3.2, 1.43, 0.65, ACCENT_BLUE)
    add_textbox(slide4, x, 3.32, 1.43, 0.4, name, font_size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

add_rect(slide4, 0.5, 4.2, 12.3, 0.8, RGBColor(0xE3, 0xF2, 0xFD))
add_textbox(slide4, 0.5, 4.3, 12.3, 0.35, '数据与存储层', font_size=14, bold=True, color=ACCENT_BLUE, align=PP_ALIGN.CENTER)
add_textbox(slide4, 0.5, 4.65, 12.3, 0.35, 'SQLite/SQLAlchemy ORM ｜ AES-256-GCM 加密 ｜ JWT 身份认证 ｜ 知识库 / 情景库', font_size=11, color=GRAY, align=PP_ALIGN.CENTER)

features = [
    ('多模态融合', '文本+URL+结构化\n知识多源输入'),
    ('实时加密', 'AES-256-GCM\nPII自动脱敏'),
    ('游戏化机制', '积分/等级/勋章\n情景模拟闯关'),
    ('知识图谱', '诈骗类型关联\n溯源推理'),
]
for i, (title, desc) in enumerate(features):
    x = 0.6 + i * 3.15
    add_card(slide4, x, 5.4, 2.95, 1.6, title, desc, icon_color=ACCENT_GOLD if i == 2 else PRIMARY, tag='核心' if i == 0 else None)

add_footer(slide4, 4)

# ============================================================
# Slide 5: 核心功能 - 智能对话 + 知识库
# ============================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide5, WHITE)
add_top_bar(slide5, '核心功能（一）：智能对话与知识库', 'INTELLIGENT CHAT & KNOWLEDGE BASE')

add_rect(slide5, 0.5, 1.8, 5.8, 5.2, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
add_rect(slide5, 0.5, 1.8, 5.8, 0.07, PRIMARY)
add_textbox(slide5, 0.8, 1.95, 5.2, 0.35, '智能对话系统', font_size=18, bold=True, color=PRIMARY)

chat_items = [
    '多轮自然对话，覆盖刷单、杀猪盘、冒充公检法等 10+ 诈骗类型',
    '适老化模式：自动检测用户年龄，切换大白话表达',
    '紧急场景响应：情绪安抚 + 止损建议 + 报警引导',
    'URL钓鱼链接一键分析，识别可疑域名',
    'PII实时检测脱敏（身份证、手机号、银行卡号）',
    '法律知识咨询：量刑标准、追赃流程、维权路径',
    '知识图谱关联查询：挖掘诈骗类型间的深层联系',
]
for i, item in enumerate(chat_items):
    add_textbox(slide5, 1.0, 2.5 + i * 0.6, 5.0, 0.5, f'▸ {item}', font_size=12, color=BLACK)

add_rect(slide5, 7.0, 1.8, 5.8, 5.2, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
add_rect(slide5, 7.0, 1.8, 5.8, 0.07, ACCENT_BLUE)
add_textbox(slide5, 7.3, 1.95, 5.2, 0.35, '结构化反诈知识库', font_size=18, bold=True, color=ACCENT_BLUE)

add_multiline(slide5, 7.5, 2.5, 5.0, 4.2, [
    ('五大常见骗局分类：', True, PRIMARY),
    ('', False, BLACK),
    ('  ● 刷单返利诈骗', False, BLACK),
    ('  ● 游戏交易诈骗', False, BLACK),
    ('  ● 冒充公检法诈骗', False, BLACK),
    ('  ● 虚假投资诈骗', False, BLACK),
    ('  ● 校园贷诈骗', False, BLACK),
    ('', False, BLACK),
    ('点击任意类型，系统自动在对话区', False, GRAY),
    ('呈现详细知识点，包括：', False, GRAY),
    ('诈骗手法拆解 → 常见话术 → 识别要点 → 防范建议', True, ACCENT_BLUE),
], font_size=12)

add_footer(slide5, 5)

# ============================================================
# Slide 6: 核心功能 - 情景模拟
# ============================================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide6, WHITE)
add_top_bar(slide6, '核心功能（二）：情景模拟训练（亮点功能）', 'SCENARIO SIMULATION TRAINING')

add_rect(slide6, 0.5, 1.8, 7.5, 5.2, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
add_rect(slide6, 0.5, 1.8, 7.5, 0.07, ACCENT_GOLD)
add_textbox(slide6, 0.8, 1.95, 7, 0.35, '沉浸式反诈训练流程', font_size=18, bold=True, color=ACCENT_GOLD)

steps = [
    ('Step 1', '选择情景', '从5个真实骗局\n场景中选择'),
    ('Step 2', 'NPC对话', '系统模拟骗子\n话术进行对话'),
    ('Step 3', '做出选择', '从多个选项中\n选择应对方式'),
    ('Step 4', '实时反馈', '即时评分 +\n安全/危险判定'),
    ('Step 5', '评估评级', '通关后获得\n总评与奖励'),
]
for i, (step, title, desc) in enumerate(steps):
    x = 0.7 + i * 1.45
    add_circle(slide6, x + 0.35, 2.5, 0.45, ACCENT_GOLD)
    add_textbox(slide6, x + 0.35, 2.57, 0.45, 0.3, str(i + 1), font_size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide6, x, 3.1, 1.2, 0.3, title, font_size=11, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
    add_textbox(slide6, x, 3.4, 1.2, 0.5, desc, font_size=10, color=GRAY, align=PP_ALIGN.CENTER)

add_textbox(slide6, 1.0, 4.4, 6.5, 0.3, '五大训练情景：', font_size=13, bold=True, color=BLACK)
scenarios_data = [
    ('刷单陷阱', '★☆☆', '刷单返利', '小额返利→连单加压→激活诈骗'),
    ('冒充公检法', '★★★', '冒充公检法', '虚假案件→安全账户→保密要求'),
    ('虚假投资理财', '★★☆', '虚假投资', '高额回报→小额试水→大额亏损'),
    ('游戏交易诈骗', '★☆☆', '游戏交易', '低价装备→平台外交易→拉黑'),
    ('校园贷危局', '★★☆', '校园贷', '低息诱惑→手续费→利滚利'),
]
for i, (name, diff, tag, flow) in enumerate(scenarios_data):
    x = 0.7 + i * 1.45
    add_rect(slide6, x, 4.8, 1.35, 2.0, LIGHT_GRAY if i % 2 == 0 else WHITE, border=RGBColor(0xDD, 0xDD, 0xDD))
    add_textbox(slide6, x + 0.1, 4.9, 1.15, 0.25, name, font_size=11, bold=True, color=BLACK, align=PP_ALIGN.CENTER)
    add_textbox(slide6, x + 0.1, 5.15, 1.15, 0.2, f'{diff} {tag}', font_size=8, color=PRIMARY, align=PP_ALIGN.CENTER)
    add_textbox(slide6, x + 0.1, 5.5, 1.15, 1.1, flow, font_size=9, color=GRAY, align=PP_ALIGN.CENTER)

add_rect(slide6, 8.5, 1.8, 4.3, 5.2, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
add_rect(slide6, 8.5, 1.8, 4.3, 0.07, GREEN)
add_textbox(slide6, 8.8, 1.95, 3.7, 0.35, '寓教于练', font_size=18, bold=True, color=GREEN)

add_multiline(slide6, 8.8, 2.5, 3.7, 4.0, [
    ('学习效果验证', True, BLACK),
    ('每个情景3个递进步骤，', False, GRAY),
    ('从诱饵到威胁，完整还原骗局链条', False, GRAY),
    ('', False, BLACK),
    ('即时反馈机制', True, BLACK),
    ('每次选择立即显示：', False, GRAY),
    ('  ✅ 安全判定 + 正向得分', False, GREEN),
    ('  ❌ 危险判定 + 负向扣分', False, PRIMARY),
    ('', False, BLACK),
    ('游戏化激励', True, BLACK),
    ('通关后根据总分评定等级：', False, GRAY),
    ('  反诈达人 · 还需警惕 · 需要提高', False, GRAY),
    ('获得积分奖励 + 成就勋章', False, ACCENT_BLUE),
], font_size=11)

add_footer(slide6, 6)

# ============================================================
# Slide 7: 核心功能 - 举报中心 + 排行榜
# ============================================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide7, WHITE)
add_top_bar(slide7, '核心功能（三）：举报中心与游戏化激励', 'REPORT CENTER & GAMIFICATION')

add_card(slide7, 0.5, 1.9, 5.8, 3.0,
         '举报中心', '用户提交可疑链接和描述 → 系统自动进行URL安全分析 → 生成分析报告 → 记录审计日志 → 联动反诈数据库', PRIMARY, '核心')

add_multiline(slide7, 0.8, 3.4, 5.2, 1.5, [
    ('举报能力：', True, BLACK),
    ('  ● 钓鱼URL自动检测与风险评级', False, GRAY),
    ('  ● 链接内容语义分析', False, GRAY),
    ('  ● 举报记录永久保存，支持追溯', False, GRAY),
    ('  ● 审计日志确保可问责性', False, GRAY),
], font_size=12)

add_card(slide7, 7.0, 1.9, 5.8, 3.0,
         '排行榜与游戏化', '积分/等级/勋章体系 → 用户排名 → 激励持续学习与举报 → 形成正向反诈社区', ACCENT_BLUE, '亮点')

add_multiline(slide7, 7.3, 3.4, 5.2, 1.5, [
    ('游戏化体系：', True, BLACK),
    ('  ● 用户等级：Lv.1~Lv.10', False, GRAY),
    ('  ● 积分获取：场景完成 + 举报 + 签到', False, GRAY),
    ('  ● 成就勋章：反诈新兵 / 火眼金睛 / 守护先锋', False, GRAY),
    ('  ● 排行榜：实时展示用户排名与成就', False, GRAY),
], font_size=12)

add_card(slide7, 0.5, 5.3, 2.8, 1.7, '历史对话', '查看所有过往交互记录\n支持回溯与复盘分析', RGBColor(0x7B, 0x1F, 0xA2))
add_card(slide7, 3.7, 5.3, 2.8, 1.7, '隐私保护', 'AES-256-GCM加密存储\nPII自动检测脱敏\n隐私状态面板可视化', GREEN)
add_card(slide7, 6.9, 5.3, 2.8, 1.7, '每日任务', '签到打卡获取积分\n每日反诈知识推送\n持续保持防骗意识', ACCENT_GOLD)
add_card(slide7, 10.1, 5.3, 2.8, 1.7, '用户中心', '个人资料与等级\n成就徽章展示\n学习数据统计', ACCENT_BLUE)

add_footer(slide7, 7)

# ============================================================
# Slide 8: 技术架构
# ============================================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide8, WHITE)
add_top_bar(slide8, '技术架构详解', 'TECHNICAL ARCHITECTURE')

add_card(slide8, 0.5, 1.9, 3.7, 2.2, '前端层', '● HTML5 / CSS3 自适应界面\n● Vanilla JS 单页应用架构\n● WebSocket 实时通信\n● 响应式设计，支持多端适配', ACCENT_BLUE, 'Frontend')
add_card(slide8, 4.6, 1.9, 3.7, 2.2, '后端引擎层 (8大引擎)', '● 意图识别引擎：NLU语义理解\n● 风险分析引擎：多维度风险评估\n● 知识检索引擎：结构化知识匹配\n● 隐私守护引擎：PII检测+AES加密\n● URL分析引擎：钓鱼链接检测\n● LLM客户端：大模型对话接入\n● 响应生成引擎：多模态输出整合\n● 知识图谱引擎：诈骗关联推理', PRIMARY, 'Backend')
add_card(slide8, 8.7, 1.9, 3.7, 2.2, '数据与安全层', '● SQLite + SQLAlchemy ORM\n● AES-256-GCM 端到端加密\n● JWT Token 身份认证\n● 审计日志完整记录\n● 隐私数据自动脱敏', GREEN, 'Data')

add_textbox(slide8, 0.8, 4.5, 12, 0.3, '技术栈明细', font_size=16, bold=True, color=BLACK)

tech_data = [
    ('后端框架', 'Python Flask 3.0', '路由管理 / RESTful API / JWT中间件'),
    ('数据库', 'SQLite + SQLAlchemy', 'ORM映射 / 迁移管理 / 连接池'),
    ('加密方案', 'AES-256-GCM', '对称加密 / 认证标签 / 密钥派生'),
    ('大语言模型', 'LLM API 接入', '多轮对话 / 意图理解 / 内容生成'),
    ('前端技术', 'HTML5 + CSS3 + JS', 'SPA / WebSocket / 响应式布局'),
    ('部署方案', 'Gunicorn / Nginx', 'WSGI服务器 / 反向代理 / 静态资源'),
]
for i, (cat, tech, desc) in enumerate(tech_data):
    y = 4.9 + i * 0.38
    col = LIGHT_GRAY if i % 2 == 0 else WHITE
    add_rect(slide8, 0.7, y, 11.5, 0.35, col)
    add_textbox(slide8, 0.9, y + 0.02, 1.8, 0.3, cat, font_size=11, bold=True, color=PRIMARY)
    add_textbox(slide8, 2.8, y + 0.02, 3.0, 0.3, tech, font_size=11, bold=True, color=BLACK)
    add_textbox(slide8, 5.9, y + 0.02, 6.0, 0.3, desc, font_size=10, color=GRAY)

add_footer(slide8, 8)

# ============================================================
# Slide 9: 创新亮点
# ============================================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide9, WHITE)
add_top_bar(slide9, '创新亮点与技术优势', 'INNOVATION & ADVANTAGES')

innovations = [
    ('多模态融合', '文本对话 + URL链接 + 结构化知识库 + 情景模拟\n四维一体，覆盖反诈全场景', '传统反诈工具功能单一，本系统实现多渠道信息融合处理'),
    ('适老化设计', '自动检测用户年龄，切换通俗语言模式\n大字体、大白话、耐心引导，降低使用门槛', '老年群体是诈骗重灾区，但多数互联网产品忽视其使用习惯'),
    ('游戏化教育', '积分/等级/勋章体系 + 情景模拟闯关\n将枯燥的反诈知识转化为沉浸式互动体验', '"看过就忘"是反诈宣传的最大痛点，互动实操才能真正建立防范意识'),
    ('隐私保护', 'AES-256-GCM加密 + PII实时脱敏\n隐私状态面板可视化，用户数据透明可控', '反诈场景涉及敏感个人信息，隐私保护既是技术底线也是信任基础'),
    ('知识图谱', '构建诈骗类型关联网络\n支持溯源推理，揭示骗术演化路径', '传统关键词匹配无法理解诈骗间的深层联系，知识图谱弥补了这一缺陷'),
    ('工程化实践', '模块化8引擎架构 + ORM数据层\n审计日志 + JWT认证 + RESTful API设计', '系统具备良好的可扩展性和可维护性，便于后续功能迭代和部署'),
]

for i, (title, desc, advantage) in enumerate(innovations):
    row = i // 2
    col = i % 2
    x = 0.5 + col * 6.2
    y = 1.85 + row * 1.8
    add_rect(slide9, x, y, 6.0, 1.65, WHITE, border=RGBColor(0xE0, 0xE0, 0xE0))
    add_rect(slide9, x, y, 0.07, 1.65, PRIMARY if i < 3 else ACCENT_BLUE)
    add_textbox(slide9, x + 0.3, y + 0.1, 5.4, 0.3, title, font_size=15, bold=True, color=PRIMARY if i < 3 else ACCENT_BLUE)
    add_textbox(slide9, x + 0.3, y + 0.45, 5.4, 0.6, desc, font_size=11, color=BLACK)
    add_rect(slide9, x + 0.3, y + 1.2, 5.4, 0.35, LIGHT_RED if i < 3 else RGBColor(0xE3, 0xF2, 0xFD))
    add_textbox(slide9, x + 0.4, y + 1.22, 5.2, 0.3, f'优势：{advantage}', font_size=9, color=PRIMARY if i < 3 else ACCENT_BLUE)

add_footer(slide9, 9)

# ============================================================
# Slide 10: 演示视频
# ============================================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide10, WHITE)
add_top_bar(slide10, '功能演示概览', 'DEMO HIGHLIGHTS')

add_textbox(slide10, 0.8, 1.8, 12, 0.4, '演示视频覆盖功能（时长 2分36秒）', font_size=16, bold=True, color=BLACK)

chapters = [
    ('聊天对话', '11轮多场景交互', '覆盖刷单识别、杀猪盘、老年人模式、紧急劝阻、URL分析、PII检测、法律咨询、知识图谱'),
    ('知识库', '3种类型深入查询', '刷单返利 → 冒充公检法 → 虚假投资，每个类型展示详细知识内容'),
    ('情景模拟', '完整3步通关', '刷单陷阱情景：从诱饵到威胁，全流程交互式训练，获得评级与奖励'),
    ('举报中心', '表单填写提交', '钓鱼URL + 描述信息 → 系统安全分析'),
    ('排行榜', '用户排名展示', '积分/等级/勋章实时排名'),
    ('隐私保护', '安全状态面板', '加密状态 + 数据保护可视化'),
]
for i, (title, duration, desc) in enumerate(chapters):
    y = 2.3 + i * 0.85
    add_rect(slide10, 0.8, y, 11.5, 0.72, LIGHT_GRAY if i % 2 == 0 else WHITE)
    add_circle(slide10, 1.0, y + 0.15, 0.4, PRIMARY if i < 3 else ACCENT_BLUE)
    add_textbox(slide10, 1.0, y + 0.18, 0.4, 0.3, str(i + 1), font_size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(slide10, 1.6, y + 0.05, 2.0, 0.3, title, font_size=15, bold=True, color=BLACK)
    add_textbox(slide10, 1.6, y + 0.35, 2.0, 0.25, duration, font_size=10, color=GRAY)
    add_textbox(slide10, 3.8, y + 0.1, 8.2, 0.5, desc, font_size=12, color=GRAY)

add_footer(slide10, 10)

# ============================================================
# Slide 11: 总结与展望
# ============================================================
slide11 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide11, WHITE)
add_top_bar(slide11, '总结与展望', 'SUMMARY & FUTURE')

add_textbox(slide11, 0.8, 1.8, 5.5, 0.35, '项目价值总结', font_size=18, bold=True, color=PRIMARY)
add_multiline(slide11, 0.8, 2.3, 5.5, 3.0, [
    ('本系统以"让反诈知识触手可及"为核心理念，', False, BLACK),
    ('构建了一套完整的多模态智能反诈对话平台：', False, BLACK),
    ('', False, BLACK),
    ('✅  10+ 诈骗类型知识覆盖', True, GREEN),
    ('✅  8大引擎协同处理', True, GREEN),
    ('✅  5个交互式情景训练', True, GREEN),
    ('✅  端到端隐私加密保护', True, GREEN),
    ('✅  游戏化激励机制', True, GREEN),
    ('✅  完整的工程化实践', True, GREEN),
], font_size=13)

add_textbox(slide11, 7.0, 1.8, 5.5, 0.35, '未来展望', font_size=18, bold=True, color=ACCENT_BLUE)
add_multiline(slide11, 7.0, 2.3, 5.5, 3.0, [
    ('短期（1-3个月）：', True, ACCENT_BLUE),
    ('  ○ 接入语音识别，支持电话诈骗实时预警', False, GRAY),
    ('  ○ 扩展知识库至30+诈骗类型', False, GRAY),
    ('  ○ 移动端PWA适配', False, GRAY),
    ('', False, BLACK),
    ('中期（3-6个月）：', True, ACCENT_BLUE),
    ('  ○ 社区联动举报网络', False, GRAY),
    ('  ○ AI驱动的诈骗手法演化追踪', False, GRAY),
    ('  ○ 多语言支持', False, GRAY),
    ('', False, BLACK),
    ('长期愿景：', True, ACCENT_BLUE),
    ('成为全民反诈基础设施，让每个人', False, GRAY),
    ('都能获得及时、专业、可信的反诈服务。', False, GRAY),
], font_size=12)

add_footer(slide11, 11)

# ============================================================
# Slide 12: 致谢
# ============================================================
slide12 = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide12, DARK_BG)
add_rect(slide12, 0, 0, 13.333, 0.15, PRIMARY)
add_rect(slide12, 0, 7.35, 13.333, 0.15, PRIMARY)

add_textbox(slide12, 0, 2.0, 13.333, 1.2, '感谢聆听', font_size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER, font_name=FONT_TITLE)
add_textbox(slide12, 0, 3.2, 13.333, 0.6, 'THANK YOU', font_size=20, color=RGBColor(0x88, 0x88, 0x88), align=PP_ALIGN.CENTER)
add_rect(slide12, 5.5, 4.0, 2.3, 0.05, ACCENT_GOLD)
add_textbox(slide12, 0, 4.3, 13.333, 0.5, '多模态智能反诈对话系统', font_size=18, color=RGBColor(0xCC, 0xCC, 0xCC), align=PP_ALIGN.CENTER)
add_textbox(slide12, 0, 4.9, 13.333, 0.5, '睿抗机器人开发者大赛 (RAICOM)', font_size=14, color=GRAY, align=PP_ALIGN.CENTER)
add_textbox(slide12, 0, 5.6, 13.333, 0.5, '用AI技术守护每一位用户的财产安全', font_size=14, color=ACCENT_GOLD, align=PP_ALIGN.CENTER)

output_path = os.path.join(BASE_DIR, '项目介绍PPT_反诈卫士_原版.pptx')
prs.save(output_path)
print(f'[OK] 原版PPT已生成: {output_path}')
print(f'[OK] 共 {len(prs.slides)} 页幻灯片')
